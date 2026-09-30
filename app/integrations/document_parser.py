"""PDF、DOCX和TXT文档解析器。"""

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pymupdf
from docx import Document
from docx.document import Document as DocxDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.core.config import settings
from app.integrations.ocr import OCRError, TesseractOCRRunner


class DocumentParseError(RuntimeError):
    """文档损坏、加密、编码未知或没有可提取文本。"""


@dataclass(frozen=True)
class ParsedSegment:
    """文档中的一个可追溯结构化文本片段。"""

    text: str
    page_number: int | None = None
    paragraph_index: int | None = None
    heading: str | None = None
    heading_level: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)


@dataclass(frozen=True)
class ParsedDocument:
    """统一的文档解析结果，同时保留兼容用全文和结构化片段。"""

    text: str
    page_count: int | None
    segments: list[ParsedSegment] = field(default_factory=list)


def _ocr_page_segments(
    runner: TesseractOCRRunner,
    page_number: int,
    text: str,
    paragraph_index: int,
) -> tuple[list[ParsedSegment], int]:
    """将单页OCR文本按非空行转换为近似段落。"""
    segments: list[ParsedSegment] = []
    for line_index, line in enumerate(text.splitlines()):
        normalized = line.strip()
        if not normalized:
            continue
        segments.append(
            ParsedSegment(
                text=normalized,
                page_number=page_number,
                paragraph_index=paragraph_index,
                metadata={
                    "source_type": "pdf",
                    "extraction_method": "ocr",
                    "ocr_engine": "tesseract",
                    "ocr_languages": runner.config.languages,
                    "ocr_line_index": line_index,
                },
            )
        )
        paragraph_index += 1
    return segments, paragraph_index


def _parse_pdf(
    path: Path,
    ocr_runner: TesseractOCRRunner | None = None,
) -> ParsedDocument:
    """逐页提取PDF文本并保留每个片段的页码。"""
    try:
        with pymupdf.open(path) as document:
            if document.needs_pass:
                raise DocumentParseError("暂不支持加密PDF")
            native_text_by_page = {
                page_number: page.get_text("text").strip()
                for page_number, page in enumerate(document, start=1)
            }
            runner = ocr_runner
            ocr_results = {}
            if settings.PDF_OCR_ENABLED:
                runner = runner or TesseractOCRRunner()
                pages_to_ocr = [
                    page_number
                    for page_number, text in native_text_by_page.items()
                    if len(text) < runner.config.min_native_text_chars
                ]
                try:
                    ocr_results = runner.extract_pages(document, pages_to_ocr)
                except OCRError as exc:
                    raise DocumentParseError(f"PDF OCR失败[{exc.code}]: {exc}") from exc

            segments: list[ParsedSegment] = []
            paragraph_index = 0
            for page_number, native_text in native_text_by_page.items():
                if page_number in ocr_results:
                    page_segments, paragraph_index = _ocr_page_segments(
                        runner,
                        page_number,
                        ocr_results[page_number].text,
                        paragraph_index,
                    )
                    segments.extend(page_segments)
                elif native_text:
                    segments.append(
                        ParsedSegment(
                            text=native_text,
                            page_number=page_number,
                            metadata={"source_type": "pdf"},
                        )
                    )
            text = "\n\n".join(segment.text for segment in segments).strip()
            return ParsedDocument(
                text=text,
                page_count=document.page_count,
                segments=segments,
            )
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(f"PDF解析失败: {exc}") from exc


def _iter_docx_blocks(document: DocxDocument):
    """按DOCX正文顺序迭代段落和表格。"""
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def _heading_level(paragraph: Paragraph) -> int | None:
    """从Word标题样式读取标题级别，无法确认时返回空值。"""
    style_name = getattr(paragraph.style, "name", "") or ""
    match = re.search(r"(?:Heading|标题)\s*(\d+)", style_name, flags=re.IGNORECASE)
    return int(match.group(1)) if match else None


def _parse_docx(path: Path) -> ParsedDocument:
    """提取DOCX段落和表格，并保留可获得的结构信息。"""
    try:
        document = Document(path)
    except Exception as exc:
        raise DocumentParseError(f"DOCX解析失败: {exc}") from exc

    segments: list[ParsedSegment] = []
    heading_path: list[str] = []
    table_index = 0
    paragraph_index = 0
    for block in _iter_docx_blocks(document):
        if isinstance(block, Paragraph):
            current_paragraph_index = paragraph_index
            paragraph_index += 1
            paragraph_text = block.text.strip()
            if not paragraph_text:
                continue
            level = _heading_level(block)
            if level is not None:
                heading_path = heading_path[: level - 1]
                heading_path.append(paragraph_text)
            active_heading = heading_path[-1] if heading_path else None
            metadata: dict[str, Any] = {
                "source_type": "docx",
                "heading_path": list(heading_path),
            }
            segments.append(
                ParsedSegment(
                    text=paragraph_text,
                    paragraph_index=current_paragraph_index,
                    heading=active_heading,
                    heading_level=level,
                    metadata=metadata,
                )
            )
            continue

        table = block
        table_index += 1
        for row_index, row in enumerate(table.rows):
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if not cells:
                continue
            segments.append(
                ParsedSegment(
                    text=" | ".join(cells),
                    heading=heading_path[-1] if heading_path else None,
                    metadata={
                        "source_type": "docx",
                        "table_index": table_index,
                        "row_index": row_index,
                        "heading_path": list(heading_path),
                    },
                )
            )

    return ParsedDocument(
        text="\n".join(segment.text for segment in segments).strip(),
        page_count=None,
        segments=segments,
    )


def _parse_txt(path: Path) -> ParsedDocument:
    """依次尝试常见中文文本编码读取TXT。"""
    raw_content = path.read_bytes()
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = raw_content.decode(encoding).strip()
            segments = [
                ParsedSegment(
                    text=line.strip(),
                    paragraph_index=paragraph_index,
                    metadata={"source_type": "txt"},
                )
                for paragraph_index, line in enumerate(text.splitlines())
                if line.strip()
            ]
            return ParsedDocument(text=text, page_count=None, segments=segments)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("TXT编码无法识别，请使用UTF-8或GB18030编码")


def parse_document(
    path: str | Path,
    *,
    ocr_runner: TesseractOCRRunner | None = None,
) -> ParsedDocument:
    """根据扩展名解析文档，并拒绝无文本内容。"""
    document_path = Path(path)
    extension = document_path.suffix.lower()
    if extension == ".pdf":
        result = _parse_pdf(document_path, ocr_runner=ocr_runner)
    elif extension == ".docx":
        result = _parse_docx(document_path)
    elif extension == ".txt":
        result = _parse_txt(document_path)
    else:
        raise DocumentParseError("不支持的文档格式")

    if not result.text:
        raise DocumentParseError("文档中没有可提取文本，扫描PDF请先进行OCR")
    return result
