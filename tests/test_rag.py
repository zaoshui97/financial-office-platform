"""企业知识库文档上传、解析和数据库记录测试。"""

from io import BytesIO

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from app.core.config import settings
from app.features.rag.chunker import chunk_document
from app.features.rag.models import DocumentChunk
from app.integrations.document_parser import ParsedDocument, ParsedSegment, parse_document


def _authorize(client: TestClient, username: str = "rag_user") -> dict[str, str]:
    """注册并登录测试用户。"""
    client.post(
        "/api/v1/auth/register",
        json={
            "username": username,
            "email": f"{username}@example.com",
            "password": "DemoPass123!",
        },
    )
    login_response = client.post(
        "/api/v1/auth/login",
        data={"username": username, "password": "DemoPass123!"},
    )
    token = login_response.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def _create_knowledge_base(client: TestClient, headers: dict[str, str]) -> int:
    """创建测试知识库并返回ID。"""
    response = client.post(
        "/api/v1/rag/knowledge-bases",
        headers=headers,
        json={"name": "企业制度库", "description": "测试文档解析"},
    )
    assert response.status_code == 201
    return response.json()["id"]


def _docx_bytes() -> bytes:
    """生成包含段落和表格的内存DOCX。"""
    document = Document()
    document.add_heading("差旅管理制度", level=1)
    document.add_paragraph("员工出差前需要提交出差申请。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "住宿标准"
    table.cell(0, 1).text = "据实报销"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_create_knowledge_base(client: TestClient) -> None:
    """登录用户可以创建和查看自己的知识库。"""
    headers = _authorize(client)
    _create_knowledge_base(client, headers)

    response = client.get("/api/v1/rag/knowledge-bases", headers=headers)

    assert response.status_code == 200
    assert response.json()[0]["name"] == "企业制度库"


def test_upload_docx_and_read_content(
    client: TestClient, db_session, monkeypatch, tmp_path
) -> None:
    """DOCX应保存、解析并将全文写入数据库。"""
    headers = _authorize(client)
    knowledge_base_id = _create_knowledge_base(client, headers)
    monkeypatch.setattr(settings, "STORAGE_DIR", tmp_path)

    upload_response = client.post(
        f"/api/v1/rag/knowledge-bases/{knowledge_base_id}/documents",
        headers=headers,
        files={
            "file": (
                "travel.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    document_id = upload_response.json()["id"]
    content_response = client.get(
        f"/api/v1/rag/documents/{document_id}/content",
        headers=headers,
    )

    assert upload_response.status_code == 201
    assert upload_response.json()["status"] == "parsed"
    assert upload_response.json()["parsed_char_count"] > 0
    assert content_response.status_code == 200
    assert "提交出差申请" in content_response.json()["parsed_text"]
    assert "住宿标准 | 据实报销" in content_response.json()["parsed_text"]
    chunks = list(
        db_session.scalars(
            select(DocumentChunk).where(DocumentChunk.document_id == document_id)
        ).all()
    )
    assert [chunk.chunk_index for chunk in chunks] == [0]
    assert chunks[0].chunk_metadata["source_type"] == "docx"
    assert chunks[0].chunk_metadata["paragraph_indexes"] == [0, 1]
    assert chunks[0].chunk_metadata["table_indexes"] == [1]
    assert len(chunks[0].content_hash) == 64


def test_upload_utf8_txt(client: TestClient, monkeypatch, tmp_path) -> None:
    """UTF-8 TXT应被解析并保存。"""
    headers = _authorize(client)
    knowledge_base_id = _create_knowledge_base(client, headers)
    monkeypatch.setattr(settings, "STORAGE_DIR", tmp_path)
    content = "信息安全制度\n员工不得泄露客户敏感信息。"

    upload_response = client.post(
        f"/api/v1/rag/knowledge-bases/{knowledge_base_id}/documents",
        headers=headers,
        files={"file": ("security.txt", content.encode("utf-8"), "text/plain")},
    )
    document_id = upload_response.json()["id"]
    content_response = client.get(
        f"/api/v1/rag/documents/{document_id}/content",
        headers=headers,
    )

    assert upload_response.status_code == 201
    assert upload_response.json()["file_type"] == "txt"
    assert content_response.json()["parsed_text"] == content


def test_upload_gb18030_txt(client: TestClient, monkeypatch, tmp_path) -> None:
    """GB18030中文TXT应被正确解码。"""
    headers = _authorize(client)
    knowledge_base_id = _create_knowledge_base(client, headers)
    monkeypatch.setattr(settings, "STORAGE_DIR", tmp_path)
    content = "财务报销制度：发票必须真实有效。"

    upload_response = client.post(
        f"/api/v1/rag/knowledge-bases/{knowledge_base_id}/documents",
        headers=headers,
        files={"file": ("finance.txt", content.encode("gb18030"), "text/plain")},
    )
    document_id = upload_response.json()["id"]
    content_response = client.get(
        f"/api/v1/rag/documents/{document_id}/content",
        headers=headers,
    )

    assert upload_response.status_code == 201
    assert content_response.json()["parsed_text"] == content


def test_parse_pdf(tmp_path) -> None:
    """PDF解析器应提取文本并记录页数。"""
    path = tmp_path / "policy.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Enterprise travel policy")
    document.save(path)
    document.close()

    parsed = parse_document(path)

    assert parsed.page_count == 1
    assert "travel policy" in parsed.text
    assert len(parsed.segments) == 1
    assert parsed.segments[0].page_number == 1


def test_parse_docx_preserves_paragraph_indexes(tmp_path) -> None:
    """DOCX解析器应保留正文段落顺序和标题元数据。"""
    path = tmp_path / "policy.docx"
    document = Document()
    document.add_paragraph("空段落")
    document.add_heading("报销制度", level=1)
    document.add_paragraph("发票必须真实有效。")
    document.save(path)

    parsed = parse_document(path)

    assert [segment.paragraph_index for segment in parsed.segments] == [0, 1, 2]
    assert parsed.segments[1].heading == "报销制度"
    assert parsed.segments[1].heading_level == 1
    assert parsed.segments[2].metadata["heading_path"] == ["报销制度"]


def test_parse_txt_generates_segments(tmp_path) -> None:
    """TXT解析器应按非空行生成有序结构化片段。"""
    path = tmp_path / "policy.txt"
    path.write_text("第一段\n\n第二段", encoding="utf-8")

    parsed = parse_document(path)

    assert [segment.text for segment in parsed.segments] == ["第一段", "第二段"]
    assert [segment.paragraph_index for segment in parsed.segments] == [0, 2]


def test_chunker_preserves_metadata_and_stable_hash() -> None:
    """Chunker应生成有序片段并为相同正文生成相同哈希。"""
    parsed = ParsedDocument(
        text="",
        page_count=2,
        segments=[
            ParsedSegment(
                text="第一页内容",
                page_number=1,
                paragraph_index=0,
                metadata={"source_type": "pdf"},
            ),
            ParsedSegment(
                text="第二页内容",
                page_number=2,
                paragraph_index=1,
                metadata={"source_type": "pdf"},
            ),
        ],
    )

    chunks = chunk_document(parsed, chunk_size=5, overlap=0)
    repeat_chunks = chunk_document(parsed, chunk_size=5, overlap=0)

    assert [chunk.chunk_index for chunk in chunks] == list(range(len(chunks)))
    assert chunks[0].page_number == 1
    assert chunks[0].metadata["page_numbers"] == [1]
    assert chunks[1].page_number == 2
    assert chunks[1].metadata["page_numbers"] == [2]
    assert [chunk.content_hash for chunk in chunks] == [
        chunk.content_hash for chunk in repeat_chunks
    ]


def test_document_chunk_index_is_unique(
    client: TestClient, db_session, monkeypatch, tmp_path
) -> None:
    """同一文档的片段序号不能重复。"""
    headers = _authorize(client, username="chunk_unique_user")
    knowledge_base_id = _create_knowledge_base(client, headers)
    monkeypatch.setattr(settings, "STORAGE_DIR", tmp_path)
    response = client.post(
        f"/api/v1/rag/knowledge-bases/{knowledge_base_id}/documents",
        headers=headers,
        files={"file": ("policy.txt", b"policy content", "text/plain")},
    )
    assert response.status_code == 201
    document_id = response.json()["id"]
    existing = db_session.scalar(
        select(DocumentChunk).where(DocumentChunk.document_id == document_id)
    )
    assert existing is not None

    duplicate = DocumentChunk(
        owner_id=existing.owner_id,
        knowledge_base_id=existing.knowledge_base_id,
        document_id=existing.document_id,
        chunk_index=existing.chunk_index,
        chunk_text="duplicate",
        chunk_metadata={},
        content_hash="0" * 64,
    )
    db_session.add(duplicate)
    with pytest.raises(IntegrityError):
        db_session.commit()


def test_empty_txt_records_failure(client: TestClient, monkeypatch, tmp_path) -> None:
    """无有效文本的文件应返回422，并保留失败记录。"""
    headers = _authorize(client)
    knowledge_base_id = _create_knowledge_base(client, headers)
    monkeypatch.setattr(settings, "STORAGE_DIR", tmp_path)

    upload_response = client.post(
        f"/api/v1/rag/knowledge-bases/{knowledge_base_id}/documents",
        headers=headers,
        files={"file": ("empty.txt", b"   \n", "text/plain")},
    )
    list_response = client.get(
        f"/api/v1/rag/knowledge-bases/{knowledge_base_id}/documents",
        headers=headers,
    )

    assert upload_response.status_code == 422
    assert list_response.status_code == 200
    assert list_response.json()[0]["status"] == "failed"


def test_reject_unsupported_file(client: TestClient) -> None:
    """非PDF、DOCX、TXT文件应被拒绝。"""
    headers = _authorize(client)
    knowledge_base_id = _create_knowledge_base(client, headers)

    response = client.post(
        f"/api/v1/rag/knowledge-bases/{knowledge_base_id}/documents",
        headers=headers,
        files={"file": ("sheet.xlsx", b"not supported", "application/octet-stream")},
    )

    assert response.status_code == 415
