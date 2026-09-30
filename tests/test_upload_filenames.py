"""上传文件名解析、规范化与安全校验测试。"""

from collections.abc import Callable
from email.header import Header
from io import BytesIO
from urllib.parse import quote

import pymupdf
import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from app.core.config import settings
from app.features.rag.models import DocumentChunk, KnowledgeDocument
from app.integrations.file_storage import UploadFilenameError, parse_upload_filename


def _authorize(client: TestClient, username: str = "filename_user") -> dict[str, str]:
    client.post(
        "/api/v1/auth/register",
        json={
            "username": username,
            "email": f"{username}@example.com",
            "password": "DemoPass123!",
        },
    )
    response = client.post(
        "/api/v1/auth/login",
        data={"username": username, "password": "DemoPass123!"},
    )
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _create_knowledge_base(client: TestClient, headers: dict[str, str]) -> int:
    response = client.post(
        "/api/v1/rag/knowledge-bases",
        headers=headers,
        json={"name": "中文文件名测试库"},
    )
    assert response.status_code == 201
    return response.json()["id"]


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("金融办公平台", level=1)
    document.add_paragraph("用于验证中文DOCX文件名。")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _pdf_bytes() -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Financial office platform")
    content = document.tobytes()
    document.close()
    return content


def _multipart_body(
    *,
    boundary: str,
    content_disposition: str,
    content_type: str,
    content: bytes,
) -> bytes:
    prefix = (
        f"--{boundary}\r\n"
        f"Content-Disposition: {content_disposition}\r\n"
        f"Content-Type: {content_type}\r\n\r\n"
    ).encode("ascii")
    return prefix + content + f"\r\n--{boundary}--\r\n".encode("ascii")


def test_upload_dotnet_encoded_chinese_docx_filename(
    client: TestClient, db_session, monkeypatch, tmp_path
) -> None:
    """应优先使用.NET同时发送的UTF-8 filename*恢复中文文件名。"""
    headers = _authorize(client)
    knowledge_base_id = _create_knowledge_base(client, headers)
    monkeypatch.setattr(settings, "STORAGE_DIR", tmp_path)
    filename = "金融企业智能办公提效平台.docx"
    encoded_word = Header(filename, "utf-8").encode()
    encoded_extended = quote(filename.encode("utf-8"), safe="")
    boundary = "----dotnet-upload-boundary"
    body = _multipart_body(
        boundary=boundary,
        content_disposition=(
            f'form-data; name="file"; filename="{encoded_word}"; '
            f"filename*=utf-8''{encoded_extended}"
        ),
        content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        content=_docx_bytes(),
    )
    headers["Content-Type"] = f"multipart/form-data; boundary={boundary}"

    response = client.post(
        f"/api/v1/rag/knowledge-bases/{knowledge_base_id}/documents",
        headers=headers,
        content=body,
    )

    assert response.status_code == 201
    assert response.json()["original_filename"] == filename
    document = db_session.get(KnowledgeDocument, response.json()["id"])
    assert document is not None
    assert document.original_filename == filename
    assert document.stored_path.endswith(".docx")
    assert filename not in document.stored_path


def test_parse_ascii_filename() -> None:
    assert parse_upload_filename("travel-policy.docx") == "travel-policy.docx"


def test_parse_rfc2047_filename_without_extended_parameter() -> None:
    filename = "差旅管理制度.docx"
    encoded_word = Header(filename, "utf-8").encode()

    result = parse_upload_filename(
        encoded_word,
        f'form-data; name="file"; filename="{encoded_word}"',
    )

    assert result == filename


def test_filename_star_has_priority_over_rfc2047_filename() -> None:
    filename = "知识库制度.docx"
    encoded_extended = quote(filename.encode("utf-8"), safe="")
    misleading = Header("wrong.txt", "utf-8").encode()

    result = parse_upload_filename(
        misleading,
        (
            f'form-data; name="file"; filename="{misleading}"; '
            f"filename*=utf-8''{encoded_extended}"
        ),
    )

    assert result == filename


@pytest.mark.parametrize(
    ("filename", "content_factory", "content_type", "expected_type"),
    [
        pytest.param(
            "金融制度.pdf",
            _pdf_bytes,
            "application/pdf",
            "pdf",
            id="pdf",
        ),
        pytest.param(
            "智能办公方案.docx",
            _docx_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
            id="docx",
        ),
        pytest.param(
            "员工手册.txt",
            lambda: "员工应遵守信息安全制度。".encode(),
            "text/plain",
            "txt",
            id="txt",
        ),
    ],
)
def test_upload_native_chinese_filenames(
    client: TestClient,
    db_session,
    monkeypatch,
    tmp_path,
    filename: str,
    content_factory: Callable[[], bytes],
    content_type: str,
    expected_type: str,
) -> None:
    headers = _authorize(client, username=f"native_{expected_type}_user")
    knowledge_base_id = _create_knowledge_base(client, headers)
    monkeypatch.setattr(settings, "STORAGE_DIR", tmp_path)

    response = client.post(
        f"/api/v1/rag/knowledge-bases/{knowledge_base_id}/documents",
        headers=headers,
        files={"file": (filename, content_factory(), content_type)},
    )

    assert response.status_code == 201
    assert response.json()["original_filename"] == filename
    assert response.json()["file_type"] == expected_type
    document = db_session.get(KnowledgeDocument, response.json()["id"])
    assert document is not None
    assert document.original_filename == filename


def test_unicode_normalization_and_uppercase_extension() -> None:
    assert parse_upload_filename("报告．ＤＯＣＸ") == "报告.DOCX"


@pytest.mark.parametrize(
    "content_disposition",
    [
        "form-data; name=file; filename*=utf-8''bad%ZZ.docx",
        "form-data; name=file; filename*=gbk''%B2%E2%CA%D4.docx",
        "form-data; name=file; filename*=utf-8''%FF.docx",
        "form-data; name=file; filename*=utf-8''%252e%252e%252fsecret.docx",
        "form-data; name=file; filename*=utf-8''%2e%2e%2fsecret.docx",
    ],
)
def test_reject_invalid_or_unsafe_extended_filename(content_disposition: str) -> None:
    with pytest.raises(UploadFilenameError):
        parse_upload_filename("safe.docx", content_disposition)


@pytest.mark.parametrize(
    "filename",
    [
        "../secret.docx",
        "..\\secret.docx",
        "/absolute.docx",
        "C:\\secret.docx",
        "encoded%2fsecret.docx",
        "encoded%ZZ.docx",
        "line\rbreak.docx",
        "line\nbreak.docx",
        "nul\x00byte.docx",
        "a" * 252 + ".txt",
    ],
)
def test_reject_unsafe_plain_filename(filename: str) -> None:
    with pytest.raises(UploadFilenameError):
        parse_upload_filename(filename)


def test_reject_malformed_or_unsupported_rfc2047_filename() -> None:
    with pytest.raises(UploadFilenameError):
        parse_upload_filename("=?utf-8?B?broken.docx")
    with pytest.raises(UploadFilenameError):
        parse_upload_filename("=?x-unsupported?B?dGVzdC5kb2N4?=")


def test_invalid_filename_does_not_leave_document_chunk_or_file(
    client: TestClient, db_session, monkeypatch, tmp_path
) -> None:
    headers = _authorize(client, username="invalid_filename_user")
    knowledge_base_id = _create_knowledge_base(client, headers)
    monkeypatch.setattr(settings, "STORAGE_DIR", tmp_path)
    boundary = "----invalid-filename-boundary"
    body = _multipart_body(
        boundary=boundary,
        content_disposition=(
            "form-data; name=\"file\"; filename=\"safe.docx\"; "
            "filename*=utf-8''%2e%2e%2fsecret.docx"
        ),
        content_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        content=_docx_bytes(),
    )
    headers["Content-Type"] = f"multipart/form-data; boundary={boundary}"

    response = client.post(
        f"/api/v1/rag/knowledge-bases/{knowledge_base_id}/documents",
        headers=headers,
        content=body,
    )

    assert response.status_code == 400
    assert db_session.scalar(select(func.count(KnowledgeDocument.id))) == 0
    assert db_session.scalar(select(func.count(DocumentChunk.id))) == 0
    assert list(tmp_path.rglob("*")) == []
